"""Excel 导入导出工具函数"""

import asyncio
from datetime import date, datetime, timedelta
from typing import Optional

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from sqlalchemy import select
from sqlalchemy.ext.asyncio import AsyncSession

from app.models.opportunity import Opportunity
from app.models.user import User

# Excel 列名到数据库字段的映射
COLUMN_MAP = {
    "序号": "serial_no",
    "项目经理": "division_member_name",
    "企业和政府": "customer_type",
    "商机情况": "opportunity_status",
    "商机时长": "opportunity_duration",
    "大单规模": "deal_scale",
    "支局": "branch_office",
    "集团商机编码": "group_code",
    "商机名称": "opportunity_name",
    "客户名称": "customer_name",
    "商机预测金额(万元)": "predicted_amount",
    "商机产品类别": "product_category",
    "产品类别": "product_category",
    "商机产品": "product_name",
    "商机创建时间": "created_date",
    "最新推进时间": "last_push_date",
    "未推进时长/天": "no_push_days",
    "未推进标记": "no_push_flag",
    "关闭商机标记": "close_flag",
    "通报标记": "notify_flag",
    "客户经理": "account_manager_name",
    "责任支局长": "branch_director_name",
    "分管领导": "leader_name",
    "商机详情": "requirement_detail",
    "商机详情（客户需求）": "requirement_detail",
    "每周商机进度更新": "weekly_progress",
    "最终签约金额": "final_amount",
    "签约合同编码": "contract_code",
    "备注": "remarks",
    "预计签约日期": "estimated_sign_date",
    "是否清单级客户": "is_list_customer",
    "是否招标": "is_bid_required",
}

# Excel 导出列定义: (列名, 列宽)
EXPORT_COLUMNS = [
    ("序号", 8),
    ("商机名称", 30),
    ("客户名称", 20),
    ("商机情况", 14),
    ("预测金额(万元)", 14),
    ("产品类别", 14),
    ("支局", 12),
    ("项目经理", 12),
    ("客户经理", 12),
    ("责任支局长", 12),
    ("预计签约日期", 14),
    ("未推进(天)", 10),
    ("优先级", 8),
    ("最新进度", 40),
]

# 需要按 real_name 匹配用户的三角色字段
ROLE_FIELDS = {
    "division_member_name": "division_member_id",
    "account_manager_name": "account_manager_id",
    "branch_director_name": "branch_director_id",
}


def _parse_date(value) -> Optional[date]:
    """解析日期：支持 Excel 序列号、字符串和 datetime 对象"""
    if value is None:
        return None
    if isinstance(value, date):
        return value
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, (int, float)):
        try:
            # Excel日期序列号从1开始（1=1900-01-01），负数或<1的值无效
            if value < 1:
                return None
            return date(1899, 12, 30) + timedelta(days=int(value))
        except (ValueError, OverflowError):
            return None
    if isinstance(value, str):
        value = value.strip()
        if not value:
            return None
        for fmt in ("%Y-%m-%d", "%Y/%m/%d", "%Y.%m.%d", "%Y年%m月%d日"):
            try:
                return datetime.strptime(value, fmt).date()
            except ValueError:
                continue
        if value.isdigit():
            try:
                num = int(value)
                if num < 1:
                    return None
                return date(1899, 12, 30) + timedelta(days=num)
            except (ValueError, OverflowError):
                return None
    return None


def _parse_boolean(value) -> bool:
    """解析布尔值"""
    if isinstance(value, bool):
        return value
    if isinstance(value, str):
        return value.strip() in ("是", "Y", "y", "yes", "1", "true", "True")
    if isinstance(value, (int, float)):
        return bool(value)
    return False


def _parse_amount(value) -> Optional[float]:
    """解析金额数值"""
    if value is None:
        return None
    if isinstance(value, (int, float)):
        return float(value)
    if isinstance(value, str):
        value = value.strip().replace(",", "").replace("，", "")
        try:
            return float(value)
        except ValueError:
            return None
    return None


def parse_excel_row(headers: list[str], row: tuple) -> dict:
    """将 Excel 一行解析为字段名->值的字典（字段名为数据库列名或 _name 后缀）"""
    data = {}
    for i, header in enumerate(headers):
        header = header.strip()
        if i >= len(row):
            break
        field = COLUMN_MAP.get(header)
        if field is None:
            continue
        value = row[i]

        # 类型转换
        if field in ("predicted_amount", "final_amount"):
            value = _parse_amount(value)
        elif field in ("created_date", "last_push_date", "estimated_sign_date"):
            value = _parse_date(value)
        elif field == "is_list_customer":
            value = _parse_boolean(value)
        elif field in (
            "no_push_days", "no_push_flag", "close_flag",
            "notify_flag", "serial_no", "opportunity_duration",
        ):
            if isinstance(value, str):
                value = value.strip()
            try:
                value = int(value) if value is not None else None
            except (ValueError, TypeError):
                value = None
        elif field in (
            "customer_type", "opportunity_status", "deal_scale",
            "branch_office", "group_code", "opportunity_name",
            "customer_name", "product_category", "product_name",
            "leader_name", "contract_code",
        ):
            # VARCHAR 字段：空值/数字0 → None，否则转为字符串
            if value is None or value == 0:
                value = None
            else:
                value = str(value).strip()
        elif field in ("requirement_detail", "weekly_progress", "remarks"):
            # TEXT 字段
            if value is None or value == 0:
                value = None
            else:
                value = str(value).strip()

        data[field] = value
    return data


async def _resolve_user_id(
    db: AsyncSession, real_name: str
) -> Optional[int]:
    """按真实姓名查找用户 ID"""
    if not real_name or not str(real_name).strip():
        return None
    name = str(real_name).strip()
    result = await db.execute(select(User).where(User.real_name == name))
    user = result.scalar_one_or_none()
    return user.id if user else None


def _load_excel_rows(file_path: str) -> tuple[list[str], list[tuple]]:
    """同步加载 Excel 文件，返回 (headers, rows)"""
    wb = load_workbook(file_path, read_only=True, data_only=True)
    ws = wb.active
    raw_rows = list(ws.iter_rows(values_only=True))
    wb.close()
    if not raw_rows:
        return [], []
    headers = [str(h) if h else "" for h in raw_rows[0]]
    return headers, raw_rows[1:]


async def import_excel(file_path: str, db: AsyncSession) -> dict:
    """导入 Excel 文件，返回导入结果统计"""
    headers, rows = await asyncio.to_thread(_load_excel_rows, file_path)
    if not rows:
        return {"total": 0, "imported": 0, "errors": ["空文件"]}

    imported = 0
    errors = []

    for row_idx, row in enumerate(rows, start=2):
        # 每行一个 savepoint，失败时只回滚当前行
        savepoint = await db.begin_nested()
        try:
            data = parse_excel_row(headers, row)
            opp_name = data.get("opportunity_name")
            if not opp_name:
                errors.append(f"第{row_idx}行：缺少商机名称，跳过")
                await savepoint.rollback()
                continue

            # 用户匹配：按 real_name 查找用户ID
            for name_field, id_field in ROLE_FIELDS.items():
                name_val = data.pop(name_field, None)
                if name_val:
                    user_id = await _resolve_user_id(db, str(name_val))
                    if user_id:
                        data[id_field] = user_id

            # 构建 Opportunity 对象（只保留有效字段）
            valid_fields = {
                k: v for k, v in data.items()
                if v is not None and k not in ROLE_FIELDS
                and hasattr(Opportunity, k)
            }

            # 自动计算 no_push_days
            if valid_fields.get("no_push_days") is None or valid_fields.get("no_push_days") == 0:
                if valid_fields.get("last_push_date"):
                    last_push = valid_fields["last_push_date"]
                    if isinstance(last_push, datetime):
                        last_push = last_push.date()
                    valid_fields["no_push_days"] = max((date.today() - last_push).days, 0)
                else:
                    valid_fields["no_push_days"] = 0

            opp = Opportunity(**valid_fields)
            db.add(opp)
            # 执行 flush 立即发现类型错误，不等到外层 flush
            try:
                await db.flush()
                await savepoint.commit()
                imported += 1
            except Exception as e:
                await savepoint.rollback()
                errors.append(f"第{row_idx}行：数据错误 - {str(e)[:100]}")
        except Exception as e:
            errors.append(f"第{row_idx}行：{str(e)[:100]}")

    return {
        "total": len(rows),
        "imported": imported,
        "errors": errors[:50],
    }


def _build_excel(opportunities: list[Opportunity], file_path: str):
    """同步构建并保存 Excel 文件"""
    wb = Workbook()
    ws = wb.active
    ws.title = "商机列表"

    # Header style
    header_font = Font(name="微软雅黑", bold=True, size=11, color="FFFFFF")
    header_fill = PatternFill(start_color="0052CC", end_color="0052CC", fill_type="solid")
    header_align = Alignment(horizontal="center", vertical="center")
    thin_border = Border(
        left=Side(style="thin", color="DFE1E6"),
        right=Side(style="thin", color="DFE1E6"),
        top=Side(style="thin", color="DFE1E6"),
        bottom=Side(style="thin", color="DFE1E6"),
    )

    # Write headers
    for col_idx, (col_name, col_width) in enumerate(EXPORT_COLUMNS, start=1):
        cell = ws.cell(row=1, column=col_idx, value=col_name)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = header_align
        cell.border = thin_border
        ws.column_dimensions[get_column_letter(col_idx)].width = col_width

    # Write data
    for row_idx, opp in enumerate(opportunities, start=2):
        row_data = [
            opp.serial_no,
            opp.opportunity_name,
            opp.customer_name,
            opp.opportunity_status,
            float(opp.predicted_amount) if opp.predicted_amount else None,
            opp.product_category,
            opp.branch_office,
            opp.division_member.real_name if getattr(opp, "division_member", None) else "",
            opp.account_manager.real_name if getattr(opp, "account_manager", None) else "",
            opp.branch_director.real_name if getattr(opp, "branch_director", None) else "",
            opp.estimated_sign_date.strftime("%Y-%m-%d") if opp.estimated_sign_date else "",
            opp.no_push_days,
            opp.priority_level,
            (opp.weekly_progress or "")[:200],
        ]
        for col_idx, value in enumerate(row_data, start=1):
            cell = ws.cell(row=row_idx, column=col_idx, value=value)
            cell.font = Font(name="微软雅黑", size=10)
            cell.border = thin_border
            if col_idx in (1, 6, 7, 12, 13):
                cell.alignment = Alignment(horizontal="center")

    wb.save(file_path)
    wb.close()


async def export_excel(opportunities: list[Opportunity], file_path: str):
    """导出商机列表到 Excel"""
    await asyncio.to_thread(_build_excel, opportunities, file_path)
